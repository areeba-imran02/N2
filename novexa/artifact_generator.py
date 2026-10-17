import io
def csv_bytes(df): return df.to_csv(index=False).encode()
def xlsx_bytes(df):
    b = io.BytesIO(); df.to_excel(b, index=False, engine="openpyxl"); return b.getvalue()
def docx_report(title, sections):
    import docx
    d = docx.Document(); d.add_heading(title, 0)
    for h, body in sections: d.add_heading(h, 1); d.add_paragraph(str(body))
    b = io.BytesIO(); d.save(b); return b.getvalue()
def txt_report(title, sections): return (title + "\n\n" + "\n\n".join(f"{h}\n{b}" for h, b in sections)).encode()
def validate(data, kind): return bool(data) and (data[:2] == b"PK" if kind in ("xlsx", "docx") else True)

def pdf_report(title, sections):
    import fitz
    doc = fitz.open(); text = title + "\n\n" + "\n\n".join(f"{h}\n{b}" for h, b in sections)
    chunks = [text[i:i + 3000] for i in range(0, len(text), 3000)] or [title]
    for c in chunks:
        page = doc.new_page(); page.insert_textbox(fitz.Rect(50, 50, 545, 790), c, fontsize=10)
    return doc.tobytes()
def make(kind, title, sections):
    if not sections or not any(str(b).strip() for _, b in sections): raise ValueError("nothing to report: no results yet")
    f = {"txt": txt_report, "docx": docx_report, "pdf": pdf_report}.get(kind)
    if not f: raise ValueError("unsupported format")
    data = f(title, sections)
    if kind == "pdf" and data[:4] != b"%PDF": raise ValueError("PDF generation failed")
    if not validate(data, kind): raise ValueError("generated file failed validation")
    return data
